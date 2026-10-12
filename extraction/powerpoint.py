"""Recursive PPTX extraction for slides, groups, tables, charts, notes, and images."""

from __future__ import annotations

from io import BytesIO
from typing import Any

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from core.document_model import DocumentBlock
from extraction.ocr import OCRStatus, configure_tesseract, ocr_image


def extract_pptx(data: bytes, ocr_status: OCRStatus | None = None) -> tuple[list[DocumentBlock], list[str]]:
    presentation = Presentation(BytesIO(data))
    blocks: list[DocumentBlock] = []
    warnings: list[str] = []
    status = ocr_status or configure_tesseract()

    def add(block_id: str, kind: str, text: str, location: str, method: str, **metadata: Any) -> None:
        if text.strip():
            blocks.append(DocumentBlock(block_id, kind, text.strip(), location, method, metadata))

    def visit_shape(shape: Any, slide_number: int, path: str) -> None:
        location = f"slide {slide_number}, shape {path} ({shape.name})"
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            for child_index, child in enumerate(shape.shapes, start=1):
                visit_shape(child, slide_number, f"{path}.{child_index}")
            return

        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            result = ocr_image(shape.image.blob, status)
            if result.text.strip():
                add(f"pptx-s{slide_number}-sh{path}-image", "image_ocr", result.text,
                    location, "ocr", slide=slide_number, shape_id=shape.shape_id,
                    word_boxes=result.boxes, ocr_engine="Tesseract")
            elif result.warning:
                warnings.append(f"{location}: {result.warning}")
            else:
                warnings.append(f"{location} produced no OCR text.")
            return

        if getattr(shape, "has_table", False):
            for row_number, row in enumerate(shape.table.rows, start=1):
                for column_number, cell in enumerate(row.cells, start=1):
                    add(f"pptx-s{slide_number}-sh{path}-r{row_number}-c{column_number}",
                        "table_cell", cell.text, f"{location}, row {row_number}, column {column_number}",
                        "native", slide=slide_number, shape_id=shape.shape_id,
                        row=row_number, column=column_number)
            return

        if getattr(shape, "has_chart", False):
            chart = shape.chart
            if chart.has_title:
                add(f"pptx-s{slide_number}-sh{path}-chart-title", "chart_text",
                    chart.chart_title.text_frame.text, f"{location}, chart title", "native",
                    slide=slide_number, shape_id=shape.shape_id)
            try:
                categories = [str(category.label) for category in chart.plots[0].categories]
                add(f"pptx-s{slide_number}-sh{path}-chart-categories", "chart_text",
                    " | ".join(categories), f"{location}, chart categories", "native",
                    slide=slide_number, shape_id=shape.shape_id)
                for series_number, series in enumerate(chart.series, start=1):
                    add(f"pptx-s{slide_number}-sh{path}-series-{series_number}", "chart_text",
                        str(series.name), f"{location}, chart series {series_number}", "native",
                        slide=slide_number, shape_id=shape.shape_id)
            except (AttributeError, IndexError, TypeError, ValueError):
                warnings.append(f"{location} contains chart data that could not be fully extracted.")
            return

        if getattr(shape, "has_text_frame", False):
            add(f"pptx-s{slide_number}-sh{path}", "shape_text", shape.text,
                location, "native", slide=slide_number, shape_id=shape.shape_id)
            return

        unsupported = {
            MSO_SHAPE_TYPE.EMBEDDED_OLE_OBJECT,
            MSO_SHAPE_TYPE.LINKED_OLE_OBJECT,
            MSO_SHAPE_TYPE.MEDIA,
        }
        if shape.shape_type in unsupported:
            warnings.append(f"{location} is an unsupported embedded object and may contain unprocessed content.")

    for slide_number, slide in enumerate(presentation.slides, start=1):
        for shape_index, shape in enumerate(slide.shapes, start=1):
            visit_shape(shape, slide_number, str(shape_index))
        if slide.has_notes_slide:
            for note_index, shape in enumerate(slide.notes_slide.notes_text_frame.paragraphs, start=1):
                add(f"pptx-s{slide_number}-note-{note_index}", "speaker_note", shape.text,
                    f"slide {slide_number}, speaker note {note_index}", "native", slide=slide_number)

    if not blocks:
        warnings.append("The PPTX produced no extractable text.")
    return blocks, warnings
