"""Ordered DOCX extraction for body content, headers, footers, and images."""

from __future__ import annotations

from io import BytesIO

from docx import Document
from docx.document import Document as DocumentObject
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph

from core.document_model import DocumentBlock
from extraction.ocr import OCRStatus, configure_tesseract, ocr_image


def _body_items(document: DocumentObject):
    for child in document.element.body.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, document)
        elif isinstance(child, CT_Tbl):
            yield Table(child, document)


def extract_docx(data: bytes, ocr_status: OCRStatus | None = None) -> tuple[list[DocumentBlock], list[str]]:
    document = Document(BytesIO(data))
    blocks: list[DocumentBlock] = []
    warnings: list[str] = []
    status = ocr_status or configure_tesseract()
    paragraph_number = table_number = 0

    for item in _body_items(document):
        if isinstance(item, Paragraph):
            paragraph_number += 1
            text = item.text.strip()
            if text:
                blocks.append(DocumentBlock(
                    f"docx-p{paragraph_number}", "paragraph", text,
                    f"paragraph {paragraph_number}", "native",
                    {"paragraph": paragraph_number, "style": item.style.name if item.style else None},
                ))
        else:
            table_number += 1
            for row_number, row in enumerate(item.rows, start=1):
                seen_cells: set[int] = set()
                for column_number, cell in enumerate(row.cells, start=1):
                    # Merged cells can be exposed repeatedly by python-docx.
                    cell_identity = id(cell._tc)
                    if cell_identity in seen_cells:
                        continue
                    seen_cells.add(cell_identity)
                    text = "\n".join(p.text.strip() for p in cell.paragraphs if p.text.strip())
                    if text:
                        location = f"table {table_number}, row {row_number}, column {column_number}"
                        blocks.append(DocumentBlock(
                            f"docx-t{table_number}-r{row_number}-c{column_number}",
                            "table_cell", text, location, "native",
                            {"table": table_number, "row": row_number, "column": column_number},
                        ))

    seen_parts: set[str] = set()
    for section_number, section in enumerate(document.sections, start=1):
        for label, part in (("header", section.header), ("footer", section.footer)):
            part_name = str(part.part.partname)
            if part_name in seen_parts:
                continue
            seen_parts.add(part_name)
            for paragraph_index, paragraph in enumerate(part.paragraphs, start=1):
                text = paragraph.text.strip()
                if text:
                    location = f"section {section_number} {label}, paragraph {paragraph_index}"
                    blocks.append(DocumentBlock(
                        f"docx-s{section_number}-{label}-p{paragraph_index}",
                        label, text, location, "native", {"section": section_number},
                    ))

    image_number = 0
    for relationship in document.part.rels.values():
        if not relationship.reltype.endswith("/image"):
            continue
        image_number += 1
        location = f"embedded image {image_number}"
        result = ocr_image(relationship.target_part.blob, status)
        if result.text.strip():
            blocks.append(DocumentBlock(
                f"docx-image-{image_number}", "image_ocr", result.text.strip(),
                location, "ocr", {"word_boxes": result.boxes, "method": "Tesseract"},
            ))
        elif result.warning:
            warnings.append(f"{location}: {result.warning}")
        else:
            warnings.append(f"{location} produced no OCR text.")

    unsupported_tags = {"object", "oleObject", "altChunk", "control"}
    found = {element.tag.split("}")[-1] for element in document.element.iter()} & unsupported_tags
    if found:
        warnings.append("Unsupported DOCX objects may contain unprocessed content: " + ", ".join(sorted(found)) + ".")
    if not blocks:
        warnings.append("The DOCX produced no extractable text.")
    return blocks, warnings

