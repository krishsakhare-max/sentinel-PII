"""Generate fail-closed, original-format sanitized document artifacts."""

from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
import re
from pathlib import Path
from typing import Any
import zipfile

import fitz
from docx import Document
from docx.document import Document as DocumentObject
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from core.document_model import DocumentBlock, DocumentResult
from core.pipeline import AnalysisResult
from extraction.extractor import extract_document
from extraction.ocr import OCRStatus
from privacy.detector import detect_pii


@dataclass(slots=True)
class SanitizedArtifact:
    """A generated file that is downloadable only after independent checks."""

    data: bytes = b""
    filename: str = ""
    export_format: str = "UNKNOWN"
    verification_status: str = "PENDING"
    preservation_status: str = "NOT GENERATED"
    downloadable: bool = False
    reason: str = "No sanitized artifact has been generated."


def _blocked(fmt: str, reason: str) -> SanitizedArtifact:
    return SanitizedArtifact(
        export_format=fmt,
        verification_status="BLOCKED",
        preservation_status="NOT VERIFIED",
        reason=reason,
    )


def _safe_filename(filename: str, fmt: str) -> str:
    stem = Path(filename).stem or "document"
    extension = {"PDF": ".pdf", "DOCX": ".docx", "PPTX": ".pptx"}[fmt]
    return f"sanitized_{stem}{extension}"


def _body_items(document: DocumentObject):
    for child in document.element.body.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, document)
        elif isinstance(child, CT_Tbl):
            yield Table(child, document)


def _docx_block_map(document: DocumentObject) -> dict[str, Any]:
    """Map the stable extraction block ids back to python-docx objects."""

    result: dict[str, Any] = {}
    paragraph_number = table_number = 0
    for item in _body_items(document):
        if isinstance(item, Paragraph):
            paragraph_number += 1
            result[f"docx-p{paragraph_number}"] = item
        else:
            table_number += 1
            for row_number, row in enumerate(item.rows, start=1):
                for column_number, cell in enumerate(row.cells, start=1):
                    result[f"docx-t{table_number}-r{row_number}-c{column_number}"] = cell

    for section_number, section in enumerate(document.sections, start=1):
        for label, part in (("header", section.header), ("footer", section.footer)):
            for paragraph_number, paragraph in enumerate(part.paragraphs, start=1):
                result[f"docx-s{section_number}-{label}-p{paragraph_number}"] = paragraph
    return result


def _write_docx(data: bytes, original: DocumentResult, sanitized: DocumentResult) -> bytes:
    document = Document(BytesIO(data))
    objects = _docx_block_map(document)
    for before, after in zip(original.blocks, sanitized.blocks):
        target = objects.get(before.block_id)
        if target is None:
            raise ValueError(f"Could not map DOCX block {before.block_id} back to the source.")
        if isinstance(target, Paragraph):
            target.text = after.text
        else:
            target.text = after.text

    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _pptx_block_map(presentation: Presentation) -> dict[str, Any]:
    result: dict[str, Any] = {}

    def visit(shape: Any, slide_number: int, path: str) -> None:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            for child_number, child in enumerate(shape.shapes, start=1):
                visit(child, slide_number, f"{path}.{child_number}")
            return
        prefix = f"pptx-s{slide_number}-sh{path}"
        result[prefix] = shape
        if getattr(shape, "has_table", False):
            for row_number, row in enumerate(shape.table.rows, start=1):
                for column_number, cell in enumerate(row.cells, start=1):
                    result[f"{prefix}-r{row_number}-c{column_number}"] = cell

    for slide_number, slide in enumerate(presentation.slides, start=1):
        for shape_number, shape in enumerate(slide.shapes, start=1):
            visit(shape, slide_number, str(shape_number))
        if slide.has_notes_slide:
            for note_number, paragraph in enumerate(slide.notes_slide.notes_text_frame.paragraphs, start=1):
                result[f"pptx-s{slide_number}-note-{note_number}"] = paragraph
    return result


def _write_pptx(data: bytes, original: DocumentResult, sanitized: DocumentResult) -> bytes:
    presentation = Presentation(BytesIO(data))
    objects = _pptx_block_map(presentation)
    for before, after in zip(original.blocks, sanitized.blocks):
        target = objects.get(before.block_id)
        if target is None:
            raise ValueError(f"Could not map PPTX block {before.block_id} back to the source.")
        if getattr(target, "has_text_frame", False):
            target.text = after.text
        elif hasattr(target, "text"):
            target.text = after.text
        else:
            raise ValueError(f"PPTX block {before.block_id} is not safely editable.")

    output = BytesIO()
    presentation.save(output)
    return output.getvalue()


def _write_pdf(data: bytes, original: DocumentResult, sanitized: DocumentResult) -> bytes:
    document = fitz.open(stream=data, filetype="pdf")
    try:
        for before, after in zip(original.blocks, sanitized.blocks):
            page_number = before.metadata.get("page")
            if not isinstance(page_number, int):
                raise ValueError(f"PDF block {before.block_id} has no page location.")
            page = document[page_number - 1]
            if before.block_type == "text":
                bbox = before.metadata.get("bbox")
                if not isinstance(bbox, list) or len(bbox) != 4:
                    raise ValueError(f"PDF block {before.block_id} has no text bounds.")
                rect = fitz.Rect(bbox)
            elif before.block_type == "image_ocr":
                boxes = before.metadata.get("image_bboxes")
                if not boxes:
                    raise ValueError(f"PDF image block {before.block_id} has no image bounds.")
                rects = [fitz.Rect(box) for box in boxes]
                rect = rects[0]
                for candidate in rects[1:]:
                    rect |= candidate
            else:
                rect = page.rect
            page.add_redact_annot(rect, text=after.text, fill=(1, 1, 1))

        for page in document:
            page.apply_redactions()
        output = BytesIO()
        document.save(output, garbage=4, clean=True, deflate=True)
        return output.getvalue()
    finally:
        document.close()


def _contains_raw_value(text: str, raw: str) -> bool:
    raw = raw.strip()
    if len(raw) < 4:
        return False
    return re.search(rf"(?<!\w){re.escape(raw)}(?!\w)", text, re.IGNORECASE) is not None


def _verify_artifact(
    data: bytes,
    filename: str,
    original: AnalysisResult,
    ocr_status: OCRStatus | None,
) -> tuple[bool, bool, str]:
    """Re-extract and independently rescan the generated artifact."""

    exported = extract_document(data, filename, ocr_status)
    if exported.extraction_warnings or not exported.extraction_complete:
        return False, False, "Generated artifact extraction was incomplete; download blocked."
    residual = detect_pii(exported)
    if residual.failed_blocks:
        return False, False, "Generated artifact could not be fully rescanned; download blocked."
    if residual.findings:
        return False, False, "Generated artifact still contains detected PII; download blocked."
    if original.detection:
        for finding in original.detection.findings:
            raw = finding.original_block_reference.text[finding.start_offset:finding.end_offset]
            if _contains_raw_value(exported.text, raw):
                return False, False, "A known original value remains in the generated artifact; download blocked."

    expected_ids = [block.block_id for block in original.document.blocks] if original.document else []
    actual_ids = [block.block_id for block in exported.blocks]
    preserved = expected_ids == actual_ids
    if not preserved:
        return False, False, "Generated artifact changed the extracted block structure; download blocked."
    return True, True, "Sanitized document generated and independently verified."


def generate_sanitized_artifact(
    data: bytes,
    analysis: AnalysisResult,
    ocr_status: OCRStatus | None = None,
) -> SanitizedArtifact:
    """Create an original-format artifact, failing closed on every uncertainty."""

    document = analysis.document
    fmt = document.file_type if document else (analysis.validation.detected_format or "UNKNOWN")
    if fmt not in {"PDF", "DOCX", "PPTX"}:
        return _blocked(fmt, "The original document format is unsupported for export.")
    if not analysis.decision or analysis.decision.decision != "APPROVED":
        return _blocked(fmt, "The security gate did not approve this document for export.")
    if not document or not analysis.redaction or not analysis.verification or not analysis.verification.passed:
        return _blocked(fmt, "The complete verified analysis is required before export.")
    if any(block.block_type == "image_ocr" for block in document.blocks) and fmt in {"DOCX", "PPTX"}:
        return _blocked(fmt, "Original embedded-image text cannot be safely rewritten in this format.")

    try:
        sanitized_document = DocumentResult(
            document.document_id, "sanitized", document.file_type,
            analysis.redaction.blocks, True, [],
        )
        if fmt == "PDF":
            exported = _write_pdf(data, document, sanitized_document)
        elif fmt == "DOCX":
            exported = _write_docx(data, document, sanitized_document)
        else:
            exported = _write_pptx(data, document, sanitized_document)
        filename = _safe_filename(analysis.filename, fmt)
        verified, preserved, reason = _verify_artifact(exported, filename, analysis, ocr_status)
        if not verified:
            return SanitizedArtifact(
                export_format=fmt, verification_status="BLOCKED",
                preservation_status="NOT VERIFIED", reason=reason,
            )
        return SanitizedArtifact(
            data=exported, filename=filename, export_format=fmt,
            verification_status="VERIFIED", preservation_status="PRESERVED",
            downloadable=True, reason=reason,
        )
    except (OSError, ValueError, RuntimeError, TypeError, KeyError, IndexError, zipfile.BadZipFile) as exc:
        return _blocked(fmt, f"Sanitized export failed closed: {type(exc).__name__}.")
