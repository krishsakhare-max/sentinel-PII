"""Deterministic, wholly synthetic adversarial document fixtures."""

from __future__ import annotations

from functools import lru_cache
from io import BytesIO
import zipfile

import fitz
from docx import Document
from docx.shared import Inches
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches as PptxInches

from attack_lab.models import AttackScenario, GroundTruthEntity
from attack_lab.scenarios import SCENARIO_METADATA, SCENARIO_ORDER


def _font(size: int) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    for path in (r"C:\Windows\Fonts\arial.ttf", r"C:\Windows\Fonts\calibri.ttf"):
        try:
            return ImageFont.truetype(path, size)
        except OSError:
            continue
    return ImageFont.load_default()


def _canonicalize_office(data: bytes) -> bytes:
    """Make synthetic OOXML bytes reproducible across Python processes."""

    source = BytesIO(data)
    output = BytesIO()
    with zipfile.ZipFile(source, "r") as incoming, zipfile.ZipFile(
        output, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9
    ) as outgoing:
        for name in sorted(incoming.namelist()):
            info = zipfile.ZipInfo(name, date_time=(2026, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            outgoing.writestr(info, incoming.read(name))
    return output.getvalue()


def _docx(paragraphs: list[str], image: bytes | None = None) -> bytes:
    document = Document()
    document.core_properties.title = "PII Sentinel synthetic attack fixture"
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if image:
        document.add_picture(BytesIO(image), width=Inches(5.8))
    output = BytesIO()
    document.save(output)
    return _canonicalize_office(output.getvalue())


def _text_image(text: str, *, width: int = 1800, height: int = 420) -> bytes:
    image = Image.new("RGB", (width, height), "white")
    ImageDraw.Draw(image).text((90, 150), text, fill="black", font=_font(64))
    output = BytesIO()
    image.save(output, format="PNG")
    return output.getvalue()


def _image_pdf(image_data: bytes) -> bytes:
    image = Image.open(BytesIO(image_data))
    pdf = fitz.open()
    page = pdf.new_page(width=image.width, height=image.height)
    page.insert_image(page.rect, stream=image_data)
    data = pdf.tobytes()
    pdf.close()
    return data


def _pptx_split_name() -> bytes:
    presentation = Presentation()
    presentation.core_properties.title = "PII Sentinel synthetic attack fixture"
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    first = slide.shapes.add_textbox(PptxInches(1), PptxInches(2), PptxInches(2), PptxInches(.8))
    first.text = "Nadia"
    last = slide.shapes.add_textbox(PptxInches(3.05), PptxInches(2), PptxInches(2), PptxInches(.8))
    last.text = "Soren"
    context = slide.shapes.add_textbox(PptxInches(1), PptxInches(3), PptxInches(5), PptxInches(.8))
    context.text = "Quarterly operating review"
    output = BytesIO()
    presentation.save(output)
    return _canonicalize_office(output.getvalue())


@lru_cache(maxsize=6)
def build_scenario(scenario_id: str) -> AttackScenario:
    if scenario_id not in SCENARIO_METADATA:
        raise KeyError(f"Unknown attack scenario: {scenario_id}")
    name, objective, document_format, difficulty = SCENARIO_METADATA[scenario_id]
    if scenario_id == "ordinary_pii":
        entities = (
            GroundTruthEntity("PERSON", "Elena Marwick", "labeled contact record"),
            GroundTruthEntity("EMAIL_ADDRESS", "elena.marwick@example.test", "labeled contact record"),
            GroundTruthEntity("PHONE_NUMBER", "+44 20 7946 0182", "labeled contact record"),
        )
        safe = ("Quarterly supplier review remains scheduled.",)
        data = _docx([f"Name: {entities[0].value}", f"Email: {entities[1].value}",
                       f"Phone: {entities[2].value}", safe[0]])
        filename = "attack-ordinary.docx"
        expected = "COMPLETE"
        constraints = ("No expected identifier may reach an approved payload.",)
        simulated = False
    elif scenario_id == "obfuscated_email":
        entities = (GroundTruthEntity("EMAIL_ADDRESS", "alex [at] example [dot] com", "obfuscated contact line"),)
        safe = ("Use the project mailbox for scheduling.",)
        data = _docx([safe[0], f"Contact: {entities[0].value}"])
        filename = "attack-obfuscated-email.docx"
        expected = "COMPLETE"
        constraints = ("Normalized obfuscated email must not reach an approved payload.",)
        simulated = False
    elif scenario_id == "split_shape_name":
        entities = (GroundTruthEntity("PERSON", "Nadia Soren", "adjacent slide shapes", ("Nadia", "Soren")),)
        safe = ("Quarterly operating review",)
        data = _pptx_split_name()
        filename = "attack-split-name.pptx"
        expected = "COMPLETE"
        constraints = ("Composite identity must be evaluated across independent blocks.",)
        simulated = False
    elif scenario_id == "scanned_phone":
        entities = (GroundTruthEntity("PHONE_NUMBER", "+61 2 9374 6812", "image-only labeled phone"),)
        safe = ("Regional support contact",)
        data = _image_pdf(_text_image(f"{safe[0]}  Phone: {entities[0].value}"))
        filename = "attack-scanned-phone.pdf"
        expected = "COMPLETE"
        constraints = ("Extraction must use OCR; exact phone must not reach approved payload.",)
        simulated = False
    elif scenario_id == "false_positive_trap":
        entities = ()
        safe = (
            "Policy date: 18 June 2026", "Policy version: 6.0", "Invoice total: $50,000",
            "Server configuration number: 1048576", "Northstar Synthetic Holdings procedure",
        )
        data = _docx(list(safe))
        filename = "attack-false-positive.docx"
        expected = "COMPLETE"
        constraints = ("Ground truth contains no PII; safe text should remain intact.",)
        simulated = False
    else:
        entities = (GroundTruthEntity("PHONE_NUMBER", "+1 415-555-0198", "embedded text image"),)
        safe = ("Synthetic image attachment",)
        data = _docx([safe[0]], _text_image(f"Phone: {entities[0].value}", width=1000, height=260))
        filename = "attack-incomplete-extraction.docx"
        expected = "PARTIAL"
        constraints = ("Simulated OCR outage must propagate PARTIAL, BLOCKED, and endpoint rejection.",)
        simulated = True
    return AttackScenario(
        scenario_id, name, objective, document_format, difficulty, filename, data,
        entities, safe, expected, constraints, simulated,
    )


def all_scenarios() -> list[AttackScenario]:
    return [build_scenario(scenario_id) for scenario_id in SCENARIO_ORDER]
