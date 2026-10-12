"""Ground-truth extraction accuracy and Phase 2.1 reliability tests."""

from io import BytesIO
import json
from pathlib import Path
import zipfile

import fitz
from docx import Document
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches

from core.extraction_quality import measure_text_quality
from extraction.extractor import extract_document
from extraction.ocr import OCRStatus


GROUND_TRUTH = json.loads((Path(__file__).parent / "fixtures" / "ground_truth.json").read_text())
OCR_UNAVAILABLE = OCRStatus(False, None, False, "Synthetic OCR outage.")


def text_image(text: str, size: tuple[int, int] = (1400, 260)) -> bytes:
    image = Image.new("RGB", size, "white")
    ImageDraw.Draw(image).text((40, 80), text, fill="black", font=ImageFont.load_default(size=58))
    output = BytesIO()
    image.save(output, format="PNG")
    return output.getvalue()


def blank_image() -> bytes:
    output = BytesIO()
    Image.new("RGB", (500, 200), "white").save(output, format="PNG")
    return output.getvalue()


def image_pdf(text: str) -> bytes:
    document = fitz.open()
    page = document.new_page(width=700, height=300)
    page.insert_image(page.rect, stream=text_image(text))
    data = document.tobytes()
    document.close()
    return data


def test_multicolumn_pdf_reading_order() -> None:
    document = fitz.open()
    page = document.new_page(width=600, height=800)
    for point, text in [((50, 100), "LEFT ONE synthetic content"), ((50, 160), "LEFT TWO synthetic content"),
                        ((330, 100), "RIGHT ONE synthetic content"), ((330, 160), "RIGHT TWO synthetic content")]:
        page.insert_text(point, text)
    result = extract_document(document.tobytes(), "columns.pdf", OCR_UNAVAILABLE)
    assert [block.text.split(" synthetic")[0] for block in result.blocks] == [
        "LEFT ONE", "LEFT TWO", "RIGHT ONE", "RIGHT TWO"
    ]


def test_scanned_pdf_known_ground_truth_and_traceability() -> None:
    truth = GROUND_TRUTH["scanned_pdf"]
    result = extract_document(image_pdf(truth["text"]), "scan.pdf")
    metrics = measure_text_quality(truth["text"], result.text)
    assert metrics.token_recall == 1.0
    assert result.blocks[0].metadata["page"] == truth["page"]
    assert result.blocks[0].extraction_status == truth["status"]
    assert result.extraction_complete


def test_pdf_with_native_and_image_text_extracts_both_once() -> None:
    document = fitz.open()
    page = document.new_page(width=700, height=500)
    page.insert_text((40, 50), "Native synthetic paragraph long enough to remain native text.")
    page.insert_image(fitz.Rect(40, 100, 660, 260), stream=text_image("IMAGE GROUND TRUTH"))
    result = extract_document(document.tobytes(), "mixed.pdf")
    assert result.text.count("Native synthetic") == 1
    assert result.text.upper().count("IMAGE GROUND TRUTH") == 1
    assert {block.extraction_status for block in result.blocks} == {"native", "ocr"}


def test_pptx_group_table_and_identifier_ground_truth() -> None:
    truth = GROUND_TRUTH["pptx"]
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    for index, text in enumerate(truth["ordered_text"]):
        group.shapes.add_textbox(Inches(1), Inches(1 + index), Inches(5), Inches(.5)).text = text
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(4.5), Inches(6), Inches(1)).table
    for row_index, row in enumerate(truth["table"]):
        for column_index, value in enumerate(row):
            table.cell(row_index, column_index).text = value
    output = BytesIO(); presentation.save(output)
    result = extract_document(output.getvalue(), "identifiers.pptx", OCR_UNAVAILABLE)
    for value in truth["ordered_text"]:
        assert value in result.text
    cells = [block for block in result.blocks if block.block_type == "table_cell"]
    assert [(b.metadata["row"], b.metadata["column"]) for b in cells] == [(1, 1), (1, 2), (2, 1), (2, 2)]
    assert all(block.metadata["slide"] == truth["slide"] for block in result.blocks)


def make_docx_with_image(image: bytes) -> bytes:
    document = Document()
    for text in GROUND_TRUTH["docx"]["ordered_paragraphs"]:
        document.add_paragraph(text)
    table = document.add_table(2, 2)
    for row_index, row in enumerate(GROUND_TRUTH["docx"]["table"]):
        for column_index, value in enumerate(row):
            table.cell(row_index, column_index).text = value
    document.add_picture(BytesIO(image), width=Inches(4))
    output = BytesIO(); document.save(output)
    return output.getvalue()


def test_docx_order_table_and_embedded_image_failure_is_partial() -> None:
    data = make_docx_with_image(blank_image())
    result = extract_document(data, "blank-image.docx")
    paragraphs = [block.text for block in result.blocks if block.block_type == "paragraph"]
    assert paragraphs == GROUND_TRUTH["docx"]["ordered_paragraphs"]
    assert not result.extraction_complete
    assert any("no OCR text" in warning for warning in result.extraction_warnings)


def test_docx_unsupported_image_encoding_is_reported() -> None:
    original = make_docx_with_image(text_image("SUPPORTED SOURCE"))
    source = zipfile.ZipFile(BytesIO(original))
    output = BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            content = source.read(info.filename)
            if info.filename.startswith("word/media/"):
                content = b"unsupported synthetic image encoding"
            target.writestr(info, content)
    result = extract_document(output.getvalue(), "unsupported.docx")
    assert not result.extraction_complete
    assert any("UnidentifiedImageError" in warning for warning in result.extraction_warnings)


def test_ocr_unavailable_and_empty_content_cannot_complete() -> None:
    result = extract_document(image_pdf("OCR REQUIRED"), "scan.pdf", OCR_UNAVAILABLE)
    assert not result.extraction_complete
    assert not result.blocks
    assert any("requires OCR" in warning for warning in result.extraction_warnings)


def test_selected_pages_demo_is_always_partial_and_traceable() -> None:
    document = fitz.open()
    for number in range(1, 4):
        page = document.new_page()
        page.insert_text((50, 80), f"Synthetic page {number} has sufficient native extraction content.")
    result = extract_document(document.tobytes(), "three-pages.pdf", OCR_UNAVAILABLE, selected_pages=[2])
    assert not result.extraction_complete
    assert {block.metadata["page"] for block in result.blocks} == {2}
    assert "page 2" in result.text.lower()
    assert any("only part" in warning for warning in result.extraction_warnings)
