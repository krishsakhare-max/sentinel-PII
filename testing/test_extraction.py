"""Synthetic tests for the Phase 2 extraction engine."""

from io import BytesIO

import fitz
from docx import Document
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches

from extraction.extractor import extract_document
from extraction.ocr import OCRStatus, configure_tesseract, ocr_image


OCR_UNAVAILABLE = OCRStatus(False, None, False, "OCR unavailable for synthetic test.")


def make_text_pdf() -> bytes:
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "Synthetic policy content for extraction testing only; no real personal data.")
    data = document.tobytes()
    document.close()
    return data


def make_blank_pdf() -> bytes:
    document = fitz.open()
    document.new_page()
    data = document.tobytes()
    document.close()
    return data


def make_docx() -> bytes:
    document = Document()
    document.add_paragraph("Synthetic training introduction")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Control"
    table.cell(0, 1).text = "Owner"
    table.cell(1, 0).text = "Synthetic review"
    table.cell(1, 1).text = "Test team"
    document.sections[0].header.paragraphs[0].text = "Synthetic header"
    document.sections[0].footer.paragraphs[0].text = "Synthetic footer"
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def make_pptx() -> bytes:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    textbox.text = "Synthetic organizational overview"
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(5), Inches(2)).table
    table.cell(0, 0).text = "Team"
    table.cell(0, 1).text = "Function"
    table.cell(1, 0).text = "Test group"
    table.cell(1, 1).text = "Quality"
    slide.notes_slide.notes_text_frame.text = "Synthetic speaker note"
    output = BytesIO()
    presentation.save(output)
    return output.getvalue()


def test_pdf_native_blocks_include_page_and_bbox() -> None:
    result = extract_document(make_text_pdf(), "synthetic.pdf", OCR_UNAVAILABLE)
    assert result.extraction_complete
    assert "Synthetic policy" in result.text
    assert result.blocks[0].source_location == "page 1"
    assert len(result.blocks[0].metadata["bbox"]) == 4


def test_blank_pdf_is_not_silently_complete_without_ocr() -> None:
    result = extract_document(make_blank_pdf(), "blank.pdf", OCR_UNAVAILABLE)
    assert not result.extraction_complete
    assert any("requires OCR" in warning for warning in result.extraction_warnings)


def test_docx_preserves_paragraph_table_header_and_footer() -> None:
    result = extract_document(make_docx(), "synthetic.docx", OCR_UNAVAILABLE)
    block_types = {block.block_type for block in result.blocks}
    assert result.extraction_complete
    assert {"paragraph", "table_cell", "header", "footer"} <= block_types
    assert "Synthetic review" in result.text


def test_pptx_extracts_shapes_tables_and_notes() -> None:
    result = extract_document(make_pptx(), "synthetic.pptx", OCR_UNAVAILABLE)
    block_types = {block.block_type for block in result.blocks}
    assert result.extraction_complete
    assert {"shape_text", "table_cell", "speaker_note"} <= block_types
    assert "Synthetic organizational overview" in result.text


def test_router_rejects_invalid_input_before_extraction() -> None:
    result = extract_document(b"not a PDF", "invalid.pdf", OCR_UNAVAILABLE)
    assert not result.extraction_complete
    assert not result.blocks
    assert result.extraction_warnings[0].startswith("Validation failed:")


def test_document_id_is_stable_and_content_based() -> None:
    data = make_text_pdf()
    first = extract_document(data, "first.pdf", OCR_UNAVAILABLE)
    second = extract_document(data, "second.pdf", OCR_UNAVAILABLE)
    assert first.document_id == second.document_id


def test_local_tesseract_and_english_are_available() -> None:
    status = configure_tesseract()
    assert status.available, status.reason
    assert status.english_available


def test_ocr_reads_synthetic_image() -> None:
    image = Image.new("RGB", (900, 220), "white")
    font = ImageFont.load_default(size=64)
    ImageDraw.Draw(image).text((40, 70), "SYNTHETIC TEST DOCUMENT", fill="black", font=font)
    result = ocr_image(image)
    assert result.warning is None
    assert "SYNTHETIC" in result.text.upper()
