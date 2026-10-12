"""Phase 5.1 orchestration, privacy-preview, and reliability coverage."""

from io import BytesIO

import fitz
import pytest
from docx import Document
from pptx import Presentation
from pptx.util import Inches

from core.document_model import DocumentBlock, DocumentResult
from core.frontend_state import (
    activate_upload, aligned_preview_blocks, analysis_cache_key, decision_presentation,
    highlighted_placeholders, highlighted_text,
)
from core.mock_ai import mock_ai_process
from core.pipeline import run_secure_analysis
from privacy.detector import detect_pii
from privacy.pii_model import PIIFinding
from privacy.redactor import redact_document
from privacy.verifier import _known_value_remains


def make_pdf(text: str = "Name: Emma Parker; email emma.parker@example.com") -> bytes:
    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((50, 80), text)
    return pdf.tobytes()


def make_docx() -> bytes:
    document = Document()
    document.add_paragraph("Employee ID: EMP-2048 and phone +1 202-555-0102")
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def make_pptx() -> bytes:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text = (
        "Contact: Anne-Marie O'Neil; email anne.oneil@example.com"
    )
    output = BytesIO()
    presentation.save(output)
    return output.getvalue()


@pytest.mark.parametrize("factory,filename", [
    (make_pdf, "one-click.pdf"),
    (make_docx, "one-click.docx"),
    (make_pptx, "one-click.pptx"),
])
def test_one_click_pipeline_cross_format(factory, filename) -> None:
    stages: list[str] = []
    result = run_secure_analysis(factory(), filename,
                                 progress_callback=lambda _done, _total, message: stages.append(message))
    assert result.completed
    assert result.decision.decision == "APPROVED"
    assert result.verification.passed
    assert result.redaction.redacted_count == len(result.detection.findings)
    assert stages[0].startswith("Validating") and stages[-1] == "Analysis complete"


def test_safe_document_is_approved_without_fabricated_findings() -> None:
    result = run_secure_analysis(make_pdf("Quarterly control framework overview."), "safe.pdf")
    assert result.completed
    assert result.detection.findings == []
    assert result.decision.decision == "APPROVED"


def test_selected_page_pipeline_is_blocked_and_cannot_reach_mock_ai() -> None:
    result = run_secure_analysis(make_pdf(), "partial.pdf", selected_pages=[1])
    assert result.completed
    assert result.decision.decision == "BLOCKED"
    assert result.decision.approved_payload is None
    with pytest.raises(PermissionError):
        mock_ai_process(result.payload, result.decision)


def test_repeated_sensitive_name_is_detected_and_redacted_everywhere() -> None:
    blocks = [
        DocumentBlock("a", "paragraph", "Name: Rowan Quill", "Paragraph 1", "native", {}),
        DocumentBlock("b", "paragraph", "rowan quill approved the record", "Paragraph 2", "native", {}),
    ]
    document = DocumentResult("repeat", "repeat.docx", "DOCX", blocks, True, [])
    detection = detect_pii(document)
    repeated = [finding for finding in detection.findings if finding.entity_type == "PERSON"]
    assert {finding.block_id for finding in repeated} == {"a", "b"}
    redaction = redact_document(document, detection.findings)
    assert all("rowan quill" not in block.text.casefold() for block in redaction.blocks)


def test_duplicate_email_phone_and_native_ocr_name_occurrences() -> None:
    blocks = [
        DocumentBlock("native", "paragraph",
                      "Name: Rowan Quill; email: team@example.com; phone: +1 202-555-0102",
                      "Paragraph 1", "native", {}),
        DocumentBlock("ocr", "image_text",
                      "ROWAN QUILL team@example.com +1 202-555-0102",
                      "Image 1", "ocr", {}),
    ]
    document = DocumentResult("duplicates", "duplicates.docx", "DOCX", blocks, True, [])
    detection = detect_pii(document)
    by_type = {entity: [f for f in detection.findings if f.entity_type == entity]
               for entity in ("PERSON", "EMAIL_ADDRESS", "PHONE_NUMBER")}
    assert all({finding.block_id for finding in findings} == {"native", "ocr"}
               for findings in by_type.values())
    redaction = redact_document(document, detection.findings)
    assert all("example.com" not in block.text.casefold() for block in redaction.blocks)
    assert all("202-555-0102" not in block.text for block in redaction.blocks)


def test_short_single_name_is_not_globally_propagated() -> None:
    blocks = [
        DocumentBlock("a", "paragraph", "Name: Al", "Paragraph 1", "native", {}),
        DocumentBlock("b", "paragraph", "AL is also a common abbreviation", "Paragraph 2", "native", {}),
    ]
    detection = detect_pii(DocumentResult("short", "short.docx", "DOCX", blocks, True, []))
    assert not any(f.detection_method == "RepeatedValuePropagation" for f in detection.findings)


def test_known_value_leakage_uses_boundaries_not_substrings() -> None:
    block = DocumentBlock("a", "paragraph", "Name: Anna", "Paragraph 1", "native", {})
    finding = PIIFinding("f", "PERSON", 6, 10, "a", None, "Paragraph 1",
                         "Synthetic", 1.0, block, "A••a")
    assert not _known_value_remains("Annabelle owns the control.", finding)
    assert _known_value_remains("Anna owns the control.", finding)


def test_preview_alignment_and_offset_integrity() -> None:
    block = DocumentBlock("source-1", "paragraph", "Email: first@example.com", "Page 1", "native", {"page": 1})
    document = DocumentResult("align", "align.pdf", "PDF", [block], True, [])
    detection = detect_pii(document)
    redaction = redact_document(document, detection.findings)
    original, sanitized = aligned_preview_blocks(document, redaction, "source-1")
    assert original.block_id == sanitized.block_id
    assert original.metadata["page"] == sanitized.metadata["page"]
    assert original.text == "Email: first@example.com"
    assert sanitized.text.startswith("Email: [EMAIL_")


def test_preview_highlighting_escapes_document_html() -> None:
    raw = "<script>alert(1)</script> first@example.com"
    start = raw.index("first@example.com")
    marked = highlighted_text(raw, [(start, len(raw))], "ps-sensitive")
    assert "<script>" not in marked
    assert "&lt;script&gt;" in marked
    assert '<mark class="ps-sensitive">first@example.com</mark>' in marked
    assert "ps-placeholder" in highlighted_placeholders("[EMAIL_1]")


def test_pending_approved_and_blocked_presentation_states() -> None:
    assert decision_presentation(None)[0] == "PENDING"
    approved = run_secure_analysis(make_pdf(), "approved.pdf").decision
    blocked = run_secure_analysis(make_pdf(), "blocked.pdf", selected_pages=[1]).decision
    assert decision_presentation(approved)[0] == "APPROVED"
    assert decision_presentation(blocked)[0] == "BLOCKED"


def test_different_upload_and_scope_never_reuse_stale_cache_key() -> None:
    first = make_pdf("First safe control summary.")
    second = make_pdf("Second safe control summary.")
    assert analysis_cache_key(first) != analysis_cache_key(second)
    assert analysis_cache_key(first) != analysis_cache_key(first, [1])


def test_upload_change_invalidates_current_session_selection() -> None:
    state = {"active_upload_digest": "old", "current_record_key": "old:full", "activity": {}}
    assert activate_upload(state, "new")
    assert state["active_upload_digest"] == "new"
    assert "current_record_key" not in state
    state["current_record_key"] = "new:full"
    assert not activate_upload(state, "new")
    assert state["current_record_key"] == "new:full"


def test_invalid_input_stops_after_validation() -> None:
    result = run_secure_analysis(b"not a document", "broken.pdf")
    assert not result.completed
    assert result.document is None
    assert result.decision is None
    assert result.errors
    assert result.stage_statuses["Validation"] == "Failed"
    assert result.stage_statuses["Extraction"] == "Pending"


def test_pipeline_reports_real_completed_and_blocked_stage_states() -> None:
    approved = run_secure_analysis(make_pdf(), "approved-stages.pdf")
    assert set(approved.stage_statuses.values()) == {"Completed"}
    blocked = run_secure_analysis(make_pdf(), "blocked-stages.pdf", selected_pages=[1])
    assert blocked.stage_statuses["Verification"] == "Completed"
    assert blocked.stage_statuses["Security gate"] == "Blocked"
