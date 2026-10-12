"""Phase 3 hybrid PII detection and exact-span evaluation tests."""

from io import BytesIO
import json
from pathlib import Path

import fitz
from docx import Document
from pptx import Presentation
from pptx.util import Inches

from core.document_model import DocumentBlock, DocumentResult
from extraction.extractor import extract_document
from privacy.detector import detect_pii
from privacy.evaluation import ExpectedFinding, evaluate_findings
from privacy.pii_model import PIIFinding
from privacy.span_resolver import resolve_findings


FIXTURE = json.loads((Path(__file__).parent / "fixtures" / "pii_ground_truth.json").read_text(encoding="utf-8"))


def fixture_document(*, complete: bool = True) -> tuple[DocumentResult, list[ExpectedFinding]]:
    blocks = []
    expected = []
    for index, item in enumerate(FIXTURE["blocks"], start=1):
        block = DocumentBlock(item["block_id"], "paragraph", item["text"], f"page {index}", "native", {"page": index})
        blocks.append(block)
        search_from = 0
        for entity, value in item["expected"]:
            start = item["text"].index(value, search_from)
            expected.append(ExpectedFinding(entity, item["block_id"], start, start + len(value)))
            search_from = start + len(value)
    return DocumentResult("fixture", "fixture.pdf", "PDF", blocks, complete, [] if complete else ["Synthetic partial extraction."]), expected


def entities_for(text: str) -> list[str]:
    block = DocumentBlock("block", "paragraph", text, "page 1", "native", {"page": 1})
    result = detect_pii(DocumentResult("doc", "test.pdf", "PDF", [block], True, []))
    return [finding.entity_type for finding in result.findings]


def test_email_phone_name_and_employee_id_detection() -> None:
    entities = entities_for("Name: Emma Parker; email emma.parker@example.com; phone +1 202-555-0101; Employee ID EMP-1042")
    assert {"PERSON", "EMAIL_ADDRESS", "PHONE_NUMBER", "EMPLOYEE_ID"} <= set(entities)


def test_international_and_hyphenated_names() -> None:
    entities = entities_for("Contact: José Álvarez\nManager: Anne-Marie O'Neil")
    assert entities.count("PERSON") == 2


def test_contextual_government_bank_and_dob_detection() -> None:
    entities = entities_for("DOB: 14/03/1992; SSN: 123-45-6789; Bank account: 12345678901")
    assert {"DATE_OF_BIRTH", "GOVERNMENT_ID", "BANK_ACCOUNT"} <= set(entities)


def test_non_pii_numbers_company_and_policy_date_are_negative() -> None:
    assert entities_for("Policy version 6.0 approved 18 June 2026; budget $50,000; Cadence Synthetic Group.") == []


def test_technical_names_and_organizational_approvers_are_not_people() -> None:
    entities = entities_for(
        "Name: Production Server\nApproved By: Change Control Board\n"
        "Database Name: Synthetic Warehouse\nManager timeout: 30"
    )
    assert "PERSON" not in entities


def test_offsets_and_source_traceability_use_original_block() -> None:
    text = "Email: trace.person@example.com"
    block = DocumentBlock("trace", "table_cell", text, "slide 3, row 2", "native", {"slide": 3})
    finding = detect_pii(DocumentResult("doc", "test.pptx", "PPTX", [block], True, [])).findings[0]
    assert text[finding.start_offset:finding.end_offset] == "trace.person@example.com"
    assert finding.block_id == "trace"
    assert finding.source_page_or_slide == 3
    assert finding.original_block_reference is block
    assert "trace.person@example.com" not in finding.masked_preview


def test_multiple_distinct_occurrences_are_preserved() -> None:
    result = entities_for("first@example.com and second@example.com")
    assert result.count("EMAIL_ADDRESS") == 2


def test_neighboring_table_label_provides_context_without_changing_offsets() -> None:
    label = DocumentBlock("label", "table_cell", "Bank account", "table 1, row 1, column 1", "native")
    value = DocumentBlock("value", "table_cell", "12345678901", "table 1, row 1, column 2", "native")
    result = detect_pii(DocumentResult("doc", "test.docx", "DOCX", [label, value], True, []))
    finding = next(item for item in result.findings if item.entity_type == "BANK_ACCOUNT")
    assert (finding.start_offset, finding.end_offset) == (0, len(value.text))
    assert finding.block_id == "value"


def test_overlap_resolver_uses_precedence_and_keeps_other_occurrences() -> None:
    block = DocumentBlock("b", "paragraph", "12345678901 xx 12345678901", "page 1", "native")
    common = dict(source_page_or_slide=1, source_location="page 1", confidence_score=.8,
                  original_block_reference=block, masked_preview="masked")
    findings = [
        PIIFinding("1", "PHONE_NUMBER", 0, 11, "b", detection_method="phone", **common),
        PIIFinding("2", "BANK_ACCOUNT", 0, 11, "b", detection_method="bank", **common),
        PIIFinding("3", "BANK_ACCOUNT", 15, 26, "b", detection_method="bank", **common),
    ]
    resolved = resolve_findings(findings)
    assert [(item.entity_type, item.start_offset) for item in resolved] == [("BANK_ACCOUNT", 0), ("BANK_ACCOUNT", 15)]


def test_exact_ground_truth_evaluation() -> None:
    document, expected = fixture_document()
    metrics = evaluate_findings(detect_pii(document).findings, expected)
    assert metrics.expected == 11
    assert metrics.precision >= .85
    assert metrics.recall >= .85
    assert metrics.f1 >= .85


def test_empty_document_warning() -> None:
    result = detect_pii(DocumentResult("empty", "empty.pdf", "PDF"))
    assert not result.findings
    assert "no extracted text" in result.warnings[0]


def test_partial_extraction_status_and_warnings_are_preserved() -> None:
    document, _ = fixture_document(complete=False)
    result = detect_pii(document)
    assert not result.extraction_complete
    assert result.extraction_warnings == ["Synthetic partial extraction."]


def test_missing_spacy_model_is_reported(monkeypatch) -> None:
    import privacy.detector as module
    monkeypatch.setattr(module, "get_analyzer", lambda: (_ for _ in ()).throw(OSError("missing model")))
    document, _ = fixture_document()
    result = module.detect_pii(document)
    assert not result.findings
    assert "initialization failed" in result.warnings[0]


def _detect_extracted(data: bytes, filename: str) -> set[str]:
    extracted = extract_document(data, filename)
    return {finding.entity_type for finding in detect_pii(extracted).findings}


def test_pdf_extraction_to_detection_integration() -> None:
    pdf = fitz.open(); page = pdf.new_page()
    page.insert_text((50, 80), "Name: Emma Parker; email: emma.parker@example.com; synthetic integration content.")
    assert {"PERSON", "EMAIL_ADDRESS"} <= _detect_extracted(pdf.tobytes(), "integration.pdf")


def test_docx_extraction_to_detection_integration() -> None:
    document = Document(); document.add_paragraph("Employee ID: EMP-2048; Phone: +1 202-555-0102")
    output = BytesIO(); document.save(output)
    assert {"EMPLOYEE_ID", "PHONE_NUMBER"} <= _detect_extracted(output.getvalue(), "integration.docx")


def test_pptx_extraction_to_detection_integration() -> None:
    presentation = Presentation(); slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text = "Contact: Anne-Marie O'Neil; Email: anne.oneil@example.com"
    output = BytesIO(); presentation.save(output)
    assert {"PERSON", "EMAIL_ADDRESS"} <= _detect_extracted(output.getvalue(), "integration.pptx")
